"""
Paso 0 — Convierte online_retail_II.xlsx (2 hojas) en un solo CSV de texto,
sin cambiar ningún valor, para que lo pueda leer el script de limpieza.

Uso:
    python scripts/00_excel_a_csv.py [ENTRADA.xlsx] [SALIDA.csv]

Por defecto:
    ENTRADA = online_retail_II.xlsx
    SALIDA  = trabajo/00_crudo/online_retail_II_crudo.csv

Qué hace (y qué NO hace):
- Lee TODAS las hojas, todas las columnas como texto (no pierde ceros ni convierte IDs a decimal).
- Agrega `hoja_origen` y `fila_excel` (fila real en Excel, 2 = primera fila de datos)
  para poder rastrear cualquier línea hasta el original.
- NO elimina duplicados, NO quita filas, NO cambia valores. El solapamiento entre hojas
  (si existe) solo se REPORTA: decidir qué hacer con él es tarea de la limpieza, con el usuario.
- Escribe UTF-8 sin BOM, separador ",", y un resumen JSON con filas por hoja y rango de fechas.
- Nunca escribe sobre el Excel original.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def a_texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        # openpyxl devuelve float para algunos enteros; se escribe sin ".0"
        return str(int(v))
    return str(v)


def main() -> int:
    raiz = Path(__file__).resolve().parent.parent
    entrada = Path(sys.argv[1]) if len(sys.argv) > 1 else raiz / "online_retail_II.xlsx"
    salida = Path(sys.argv[2]) if len(sys.argv) > 2 else raiz / "trabajo" / "00_crudo" / "online_retail_II_crudo.csv"
    salida.parent.mkdir(parents=True, exist_ok=True)

    if salida.resolve() == entrada.resolve():
        print("ERROR: la salida no puede ser el archivo original.")
        return 1

    wb = load_workbook(entrada, read_only=True, data_only=True)
    partes, resumen = [], {"archivo": str(entrada.name), "generado": datetime.now().isoformat(timespec="seconds"), "hojas": []}
    encabezado_ref = None

    for ws in wb.worksheets:
        filas = ws.iter_rows(values_only=True)
        encabezado = [a_texto(c).strip() for c in next(filas)]
        if encabezado_ref is None:
            encabezado_ref = encabezado
        elif encabezado != encabezado_ref:
            print(f"ERROR: la hoja '{ws.title}' tiene encabezados distintos: {encabezado} vs {encabezado_ref}")
            return 1
        datos = [[a_texto(c) for c in fila] for fila in filas]
        df = pd.DataFrame(datos, columns=encabezado, dtype=str)
        df.insert(0, "fila_excel", range(2, len(df) + 2))
        df.insert(0, "hoja_origen", ws.title)
        fechas = pd.to_datetime(df["InvoiceDate"], format="%Y-%m-%d %H:%M:%S", errors="coerce")
        resumen["hojas"].append({
            "hoja": ws.title,
            "filas": int(len(df)),
            "fecha_min": str(fechas.min()),
            "fecha_max": str(fechas.max()),
            "fechas_no_parseadas": int(fechas.isna().sum()),
        })
        partes.append(df)
        print(f"  {ws.title}: {len(df):,} filas  ({fechas.min()} -> {fechas.max()})")

    total = pd.concat(partes, ignore_index=True)

    # Solo REPORTAR el solapamiento entre hojas; no se toca nada.
    if len(partes) == 2:
        cols = encabezado_ref
        a = partes[0][cols].assign(_n=partes[0].groupby(cols).cumcount())
        b = partes[1][cols].assign(_n=partes[1].groupby(cols).cumcount())
        comunes = a.merge(b, on=cols + ["_n"], how="inner")
        f0 = pd.to_datetime(partes[0]["InvoiceDate"], errors="coerce")
        f1 = pd.to_datetime(partes[1]["InvoiceDate"], errors="coerce")
        ini, fin = max(f0.min(), f1.min()), min(f0.max(), f1.max())
        resumen["solapamiento"] = {
            "rango_fechas_comun": [str(ini), str(fin)] if ini <= fin else None,
            "filas_hoja1_en_rango": int(((f0 >= ini) & (f0 <= fin)).sum()) if ini <= fin else 0,
            "filas_hoja2_en_rango": int(((f1 >= ini) & (f1 <= fin)).sum()) if ini <= fin else 0,
            "filas_identicas_en_ambas_hojas": int(len(comunes)),
            "nota": "Solo diagnóstico. Qué hacer con estas filas se decide en la limpieza, con el usuario.",
        }

    resumen["filas_total"] = int(len(total))
    total.to_csv(salida, index=False, encoding="utf-8")
    (salida.parent / "resumen_conversion.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8")

    # Verificación: releer y contar
    releido = pd.read_csv(salida, dtype=str, keep_default_na=False)
    esperado = sum(h["filas"] for h in resumen["hojas"])
    if len(releido) != esperado:
        print(f"ERROR: se esperaban {esperado:,} filas y el CSV tiene {len(releido):,}")
        return 1
    print(f"OK: {len(releido):,} filas escritas en {salida}")
    if "solapamiento" in resumen:
        print("Solapamiento entre hojas:", json.dumps(resumen["solapamiento"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
